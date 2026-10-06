"""Append-only Phase 9 case, run, and reviewer data stores."""

from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path
from pathlib import PurePosixPath
from typing import Any, Dict, List, Optional, Type, TypeVar
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from pydantic import Field, field_validator

from trust_safety_agent.schema import StrictModel, StringEnum


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def validate_media_reference(value: str) -> str:
    cleaned = value.strip().replace("\\", "/")
    if cleaned.startswith(("http://", "https://")):
        return cleaned
    if "://" in cleaned:
        raise ValueError("media reference must use http, https, or a project-relative path")
    path = PurePosixPath(cleaned)
    if path.is_absolute() or ".." in path.parts:
        raise ValueError("media reference must be a safe project-relative path")
    return cleaned


class CaseSource(StringEnum):
    MANUAL = "manual"
    CSV = "csv"
    EXCEL = "excel"
    SYNTHETIC = "synthetic"


class ShopCase(StrictModel):
    shop_id: str = Field(min_length=1, max_length=100)
    shop_name: str = Field(min_length=1, max_length=300)
    shop_avatar: str = Field(min_length=1, max_length=1000)
    brand: Optional[str] = Field(default=None, max_length=100)
    brand_authorized: Optional[bool] = None
    source: CaseSource = CaseSource.MANUAL
    revision: int = Field(default=1, ge=1)
    created_at: datetime = Field(default_factory=utc_now)
    updated_at: datetime = Field(default_factory=utc_now)
    archived: bool = False

    @field_validator("shop_avatar")
    @classmethod
    def validate_avatar(cls, value: str) -> str:
        return validate_media_reference(value)


class ProductCase(StrictModel):
    product_id: str = Field(min_length=1, max_length=100)
    title: str = Field(min_length=1, max_length=500)
    description: str = Field(default="", max_length=4000)
    product_images: List[str] = Field(min_length=1, max_length=12)
    price: Optional[float] = Field(default=None, gt=0)
    currency: Optional[str] = Field(default=None, min_length=3, max_length=3)
    brand: Optional[str] = Field(default=None, max_length=100)
    brand_authorized: Optional[bool] = None
    source: CaseSource = CaseSource.MANUAL
    revision: int = Field(default=1, ge=1)
    created_at: datetime = Field(default_factory=utc_now)
    updated_at: datetime = Field(default_factory=utc_now)
    archived: bool = False

    @field_validator("product_images")
    @classmethod
    def validate_images(cls, images: List[str]) -> List[str]:
        if any(not image.strip() for image in images):
            raise ValueError("product_images cannot contain empty values")
        return [validate_media_reference(image) for image in images]


class LabelStatus(StringEnum):
    HIT = "hit"
    NOT_HIT = "not_hit"
    INSUFFICIENT_EVIDENCE = "insufficient_evidence"
    NOT_APPLICABLE = "not_applicable"


class RunDecision(StringEnum):
    APPROVE = "approve"
    REJECT = "reject"
    MANUAL_REVIEW = "manual_review"


class AgentRunRecord(StrictModel):
    run_id: str = Field(default_factory=lambda: f"RUN-{uuid4().hex[:12].upper()}")
    case_id: str
    case_revision: int = Field(ge=1)
    workflow_version: str
    prompt_version: str
    skill_version: str
    model_id: str
    policy_index_version: str
    brand_index_version: str
    suggested_labels: Dict[str, LabelStatus] = Field(default_factory=dict)
    suggested_decision: RunDecision
    confidence: float = Field(ge=0, le=1)
    trace_id: str
    experimental: bool = False
    parent_run_id: Optional[str] = None
    created_at: datetime = Field(default_factory=utc_now)


class ReviewerAnnotationRecord(StrictModel):
    review_id: str = Field(default_factory=lambda: f"REV9-{uuid4().hex[:12].upper()}")
    run_id: str
    reviewer_labels: Dict[str, LabelStatus] = Field(default_factory=dict)
    reviewer_decision: RunDecision
    reviewer_note: str = Field(default="", max_length=4000)
    failure_feedback: Optional[str] = Field(default=None, max_length=100)
    created_at: datetime = Field(default_factory=utc_now)


T = TypeVar("T", bound=StrictModel)


class JsonlStore:
    def __init__(self, path: Path, model: Type[T]) -> None:
        self.path = path
        self.model = model

    def append(self, record: T) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.path.open("a", encoding="utf-8") as output:
            output.write(record.model_dump_json() + "\n")

    def list(self) -> List[T]:
        if not self.path.exists():
            return []
        return [
            self.model.model_validate_json(line)
            for line in self.path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]


class CaseStore:
    """Stores immutable revisions; list methods return only the latest revision."""

    def __init__(self, root: Path) -> None:
        self.root = root
        self.shop_store = JsonlStore(root / "shop_cases.jsonl", ShopCase)
        self.product_store = JsonlStore(root / "product_cases.jsonl", ProductCase)

    @staticmethod
    def _latest(records: List[T], id_field: str) -> List[T]:
        latest: Dict[str, T] = {}
        for record in records:
            key = str(getattr(record, id_field))
            if key not in latest or record.revision > latest[key].revision:
                latest[key] = record
        return sorted(latest.values(), key=lambda item: getattr(item, id_field))

    def list_shops(self, include_archived: bool = False) -> List[ShopCase]:
        records = self._latest(self.shop_store.list(), "shop_id")
        return records if include_archived else [item for item in records if not item.archived]

    def list_products(self, include_archived: bool = False) -> List[ProductCase]:
        records = self._latest(self.product_store.list(), "product_id")
        return records if include_archived else [item for item in records if not item.archived]

    def get_shop(self, shop_id: str) -> Optional[ShopCase]:
        return next((item for item in self.list_shops(True) if item.shop_id == shop_id), None)

    def get_product(self, product_id: str) -> Optional[ProductCase]:
        return next((item for item in self.list_products(True) if item.product_id == product_id), None)

    def create_shop(self, case: ShopCase) -> ShopCase:
        if self.get_shop(case.shop_id):
            raise ValueError(f"shop_id {case.shop_id} already exists")
        self.shop_store.append(case)
        return case

    def create_product(self, case: ProductCase) -> ProductCase:
        if self.get_product(case.product_id):
            raise ValueError(f"product_id {case.product_id} already exists")
        self.product_store.append(case)
        return case

    def revise_shop(self, shop_id: str, changes: Dict[str, Any]) -> ShopCase:
        current = self.get_shop(shop_id)
        if current is None:
            raise KeyError(shop_id)
        blocked = {"shop_id", "revision", "created_at", "updated_at"}
        if blocked.intersection(changes):
            raise ValueError("identity and revision fields cannot be changed")
        revised = current.model_copy(update={**changes, "revision": current.revision + 1, "updated_at": utc_now()})
        revised = ShopCase.model_validate(revised.model_dump())
        self.shop_store.append(revised)
        return revised

    def revise_product(self, product_id: str, changes: Dict[str, Any]) -> ProductCase:
        current = self.get_product(product_id)
        if current is None:
            raise KeyError(product_id)
        blocked = {"product_id", "revision", "created_at", "updated_at"}
        if blocked.intersection(changes):
            raise ValueError("identity and revision fields cannot be changed")
        revised = current.model_copy(update={**changes, "revision": current.revision + 1, "updated_at": utc_now()})
        revised = ProductCase.model_validate(revised.model_dump())
        self.product_store.append(revised)
        return revised

    def archive_shop(self, shop_id: str) -> ShopCase:
        return self.revise_shop(shop_id, {"archived": True})

    def archive_product(self, product_id: str) -> ProductCase:
        return self.revise_product(product_id, {"archived": True})


def export_records(records: List[StrictModel], path: Path) -> None:
    if not records:
        raise ValueError("cannot export an empty record list")
    rows = [record.model_dump(mode="json") for record in records]
    for row in rows:
        for key, value in list(row.items()):
            if isinstance(value, list):
                row[key] = json.dumps(value, ensure_ascii=False)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as output:
            writer = csv.DictWriter(output, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        return
    if path.suffix.lower() != ".xlsx":
        raise ValueError("export path must use .csv or .xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(rows[0]))
    for row in rows:
        sheet.append([row[key] for key in rows[0]])
    workbook.save(path)


def read_tabular_rows(path: Path) -> List[Dict[str, Any]]:
    if not path.exists() or not path.stat().st_size:
        raise ValueError("input file is empty or missing")
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as source:
            return list(csv.DictReader(source))
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        if not rows:
            raise ValueError("input file is empty")
        return [dict(zip(rows[0], row)) for row in rows[1:]]
    raise ValueError("input path must use .csv or .xlsx")


def _nullable_bool(value: Any) -> Optional[bool]:
    if value is None or str(value).strip().casefold() in {"", "null", "none", "unknown"}:
        return None
    normalized = str(value).strip().casefold()
    if normalized in {"true", "1", "yes"}:
        return True
    if normalized in {"false", "0", "no"}:
        return False
    raise ValueError(f"invalid nullable boolean: {value}")


def import_shop_cases(path: Path, store: CaseStore) -> List[ShopCase]:
    source = CaseSource.CSV if path.suffix.lower() == ".csv" else CaseSource.EXCEL
    imported = []
    seen = set()
    for row in read_tabular_rows(path):
        shop_id = str(row.get("shop_id") or "").strip()
        if shop_id in seen:
            raise ValueError(f"duplicate shop_id in import: {shop_id}")
        seen.add(shop_id)
        case = ShopCase(
            shop_id=shop_id,
            shop_name=str(row.get("shop_name") or ""),
            shop_avatar=str(row.get("shop_avatar") or ""),
            brand=str(row.get("brand") or "").strip() or None,
            brand_authorized=_nullable_bool(row.get("brand_authorized")),
            source=source,
        )
        store.create_shop(case)
        imported.append(case)
    return imported


def import_product_cases(path: Path, store: CaseStore) -> List[ProductCase]:
    source = CaseSource.CSV if path.suffix.lower() == ".csv" else CaseSource.EXCEL
    imported = []
    seen = set()
    for row in read_tabular_rows(path):
        product_id = str(row.get("product_id") or "").strip()
        if product_id in seen:
            raise ValueError(f"duplicate product_id in import: {product_id}")
        seen.add(product_id)
        raw_images = row.get("product_images") or ""
        if isinstance(raw_images, list):
            images = raw_images
        else:
            raw = str(raw_images).strip()
            try:
                images = json.loads(raw) if raw.startswith("[") else [item.strip() for item in raw.split("|") if item.strip()]
            except json.JSONDecodeError as exc:
                raise ValueError(f"invalid product_images JSON for {product_id}") from exc
        price_value = row.get("price")
        case = ProductCase(
            product_id=product_id,
            title=str(row.get("title") or ""),
            description=str(row.get("description") or ""),
            product_images=images,
            price=float(price_value) if price_value not in (None, "") else None,
            currency=str(row.get("currency") or "").strip() or None,
            brand=str(row.get("brand") or "").strip() or None,
            brand_authorized=_nullable_bool(row.get("brand_authorized")),
            source=source,
        )
        store.create_product(case)
        imported.append(case)
    return imported

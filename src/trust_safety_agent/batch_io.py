"""Validated CSV/XLSX batch import and export for review workflows."""

from __future__ import annotations

import csv
import io
import json
from typing import Dict, Iterable, List, Sequence, TypeVar

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ValidationError

from trust_safety_agent.product_review import ProductReviewInput
from trust_safety_agent.shop_identity import ShopIdentityInput


class BatchValidationError(ValueError):
    def __init__(self, errors: Sequence[str]) -> None:
        self.errors = list(errors)
        super().__init__("; ".join(self.errors))


def _read_rows(payload: bytes, filename: str) -> tuple[List[str], List[Dict[str, str]]]:
    if not payload:
        raise BatchValidationError(["文件为空。"])
    lowered = filename.casefold()
    if lowered.endswith(".csv"):
        try:
            text = payload.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise BatchValidationError(["CSV 必须使用 UTF-8 编码。"]) from exc
        reader = csv.DictReader(io.StringIO(text))
        headers = [header.strip() for header in (reader.fieldnames or []) if header]
        rows = [
            {str(key).strip(): (value or "").strip() for key, value in row.items() if key}
            for row in reader
            if any((value or "").strip() for value in row.values())
        ]
        return headers, rows
    if lowered.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True))
        except Exception as exc:
            raise BatchValidationError(["Excel 文件无法解析。"]) from exc
        if not values:
            raise BatchValidationError(["文件为空。"])
        headers = [str(value or "").strip() for value in values[0]]
        rows = []
        for values_row in values[1:]:
            row = {
                header: str(value).strip() if value is not None else ""
                for header, value in zip(headers, values_row)
                if header
            }
            if any(row.values()):
                rows.append(row)
        return headers, rows
    raise BatchValidationError(["仅支持 .csv 和 .xlsx 文件。"])


def _validate_table(
    headers: Sequence[str],
    rows: Sequence[Dict[str, str]],
    required: Sequence[str],
) -> None:
    missing = [field for field in required if field not in headers]
    errors = [f"缺少必填列：{field}" for field in missing]
    if not rows:
        errors.append("文件没有可处理的数据行。")
    case_ids = [row.get("case_id", "") for row in rows if row.get("case_id")]
    duplicates = sorted({value for value in case_ids if case_ids.count(value) > 1})
    if duplicates:
        errors.append(f"case_id 重复：{', '.join(duplicates)}")
    if errors:
        raise BatchValidationError(errors)


def parse_product_batch(payload: bytes, filename: str) -> List[ProductReviewInput]:
    headers, rows = _read_rows(payload, filename)
    required = ["case_id", "product_id", "title", "description"]
    _validate_table(headers, rows, required)
    records = []
    errors = []
    allowed = set(ProductReviewInput.model_fields)
    for index, row in enumerate(rows, start=2):
        values = {key: value for key, value in row.items() if key in allowed}
        for optional in ("product_id", "optional_brand_field", "optional_category", "image_url", "image_path"):
            if optional in values and not values[optional]:
                values[optional] = None
        if "visual_marks" in values:
            values["visual_marks"] = [
                item.strip() for item in values["visual_marks"].split("|") if item.strip()
            ]
        try:
            records.append(ProductReviewInput.model_validate(values))
        except ValidationError as exc:
            detail = exc.errors()[0]
            field = ".".join(str(part) for part in detail["loc"])
            errors.append(f"第 {index} 行 {field}：{detail['msg']}")
    if errors:
        raise BatchValidationError(errors)
    return records


def parse_shop_batch(payload: bytes, filename: str) -> List[ShopIdentityInput]:
    headers, rows = _read_rows(payload, filename)
    required = ["case_id", "shop_name", "avatar_url", "controlled_brand"]
    _validate_table(headers, rows, required)
    records = []
    errors = []
    allowed = set(ShopIdentityInput.model_fields)
    for index, row in enumerate(rows, start=2):
        values = {key: value for key, value in row.items() if key in allowed}
        if not values.get("avatar_url"):
            values["avatar_url"] = None
        if not values.get("authorization_status"):
            values["authorization_status"] = "unknown"
        if "avatar_visual_marks" in values:
            values["avatar_visual_marks"] = [
                item.strip()
                for item in values["avatar_visual_marks"].split("|")
                if item.strip()
            ]
        try:
            records.append(ShopIdentityInput.model_validate(values))
        except ValidationError as exc:
            detail = exc.errors()[0]
            field = ".".join(str(part) for part in detail["loc"])
            errors.append(f"第 {index} 行 {field}：{detail['msg']}")
    if errors:
        raise BatchValidationError(errors)
    return records


RecordT = TypeVar("RecordT", bound=BaseModel)


def _flat_rows(records: Iterable[RecordT]) -> List[Dict[str, object]]:
    rows = []
    for record in records:
        values = record.model_dump(mode="json")
        rows.append(
            {
                key: (
                    json.dumps(value, ensure_ascii=False, separators=(",", ":"))
                    if isinstance(value, (dict, list))
                    else value
                )
                for key, value in values.items()
            }
        )
    return rows


def export_csv(records: Iterable[RecordT]) -> bytes:
    rows = _flat_rows(records)
    if not rows:
        return b""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def export_xlsx(records: Iterable[RecordT]) -> bytes:
    rows = _flat_rows(records)
    if not rows:
        return b""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "review_results"
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
